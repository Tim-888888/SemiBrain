"""Format adapters retain source coordinates without executing embedded content."""

import base64
import io
import json
import posixpath
import zipfile
from urllib.parse import unquote
from xml.etree import ElementTree as ET

from PIL import Image


def output(parts, blocks, *, images=None, findings=None, **metadata):
    return {"content": "\n\n".join(parts), "blocks": blocks, "images": images or {},
            "metadata": {**metadata, "quality_findings": findings or []}}


def picture(raw, images):
    from hashlib import sha256
    with Image.open(io.BytesIO(raw)) as image:
        if image.width * image.height > 16_000_000:
            raise ValueError("IMAGE_DIMENSION_LIMIT")
        image.verify()
        ext = {"PNG": "png", "JPEG": "jpg", "WEBP": "webp", "GIF": "gif"}.get(image.format)
    if not ext:
        raise ValueError("IMAGE_FORMAT_UNSUPPORTED")
    key = "images/" + sha256(raw).hexdigest()[:24] + "." + ext
    images[key] = base64.b64encode(raw).decode()
    return "![原图](" + key + ")"


def block(text, kind, **location):
    return {"text": text, "kind": kind, "location": {"source": "original", **location}}


def parse_json(content):
    def pairs(values):
        result = {}
        for key, value in values:
            if key in result:
                raise ValueError("JSON_DUPLICATE_KEY")
            result[key] = value
        return result

    def invalid_constant(value):
        raise ValueError("JSON_NONFINITE_VALUE")

    value = json.loads(content, object_pairs_hook=pairs, parse_constant=invalid_constant)
    blocks, parts = [], []

    def walk(item, path, depth):
        if depth > 64 or len(blocks) >= 10000:
            raise ValueError("JSON_STRUCTURE_LIMIT")
        if isinstance(item, (dict, list)) and item:
            for key, child in (item.items() if isinstance(item, dict) else enumerate(item)):
                escaped = str(key).replace("~", "~0").replace("/", "~1")
                walk(child, path + "/" + escaped, depth + 1)
        else:
            text = json.dumps(item, ensure_ascii=False)
            rendered = "路径 " + (path or "/") + "\n\n```json\n" + text + "\n```"
            parts.append(rendered)
            blocks.append(block(rendered, "json_value", json_pointer=path))

    walk(value, "", 0)
    return output(parts, blocks)


def parse_xlsx(content):
    import openpyxl
    from docreader.parser.xlsx_repair import repair_xlsx_bytes, sanitize_xlsx_styles

    findings, repaired = [], repair_xlsx_bytes(content)
    if repaired is not None:
        content = repaired
        findings.append("XLSX_PACKAGING_REPAIRED")
    try:
        formulas = openpyxl.load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    except TypeError:
        repaired = sanitize_xlsx_styles(content)
        if repaired is None:
            raise
        content = repaired
        formulas = openpyxl.load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
        findings.append("XLSX_STYLES_REPAIRED")
    cached = openpyxl.load_workbook(io.BytesIO(content), data_only=True, keep_links=False)
    parts, blocks, images, count = [], [], {}, 0
    try:
        for sheet in formulas:
            if sheet.max_row * sheet.max_column > 200000 or sheet.max_column > 512:
                raise ValueError("SPREADSHEET_CELL_LIMIT")
            header = "## 工作表：" + sheet.title
            parts.append(header)
            blocks.append(block(header, "sheet", sheet=sheet.title))
            ranges = list(sheet.merged_cells.ranges)
            for row in sheet.iter_rows():
                cells = []
                for cell in row:
                    if cell.value is None:
                        continue
                    count += 1
                    if count > 100000:
                        raise ValueError("SPREADSHEET_CELL_LIMIT")
                    value = str(cell.value)
                    location = {"sheet": sheet.title, "cell": cell.coordinate,
                                "row": cell.row, "column": cell.column}
                    merged = next((str(r) for r in ranges if cell.coordinate in r), None)
                    if merged:
                        location["merged_range"] = merged
                        value += "（合并区域 " + merged + "）"
                    if cell.data_type == "f":
                        cache = cached[sheet.title][cell.coordinate].value
                        location.update(formula=str(cell.value), cached_value=cache)
                        value += "；缓存值：" + (str(cache) if cache is not None else "缺失，未计算")
                        if cache is None:
                            findings.append("FORMULA_CACHE_MISSING")
                    text = cell.coordinate + ": " + value
                    blocks.append(block(text, "cell", **location))
                    cells.append(text)
                if cells:
                    parts.append("\n".join(cells))
            for image in sheet._images:
                try:
                    text = picture(image._data(), images)
                    anchor = image.anchor._from
                    parts.append(text)
                    blocks.append(block(text, "image", sheet=sheet.title,
                                        row=anchor.row + 1, column=anchor.col + 1))
                except (ValueError, OSError, AttributeError):
                    findings.append("EMBEDDED_IMAGE_UNSUPPORTED")
    finally:
        formulas.close()
        cached.close()
    return output(parts, blocks, images=images, findings=sorted(set(findings)))


def parse_xls(content):
    import xlrd
    from openpyxl.utils import get_column_letter
    book = xlrd.open_workbook(file_contents=content, formatting_info=True, on_demand=True)
    parts, blocks = [], []
    try:
        for sheet in book.sheets():
            if sheet.nrows * sheet.ncols > 200000:
                raise ValueError("SPREADSHEET_CELL_LIMIT")
            parts.append("## 工作表：" + sheet.name)
            for r in range(sheet.nrows):
                cells = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        continue
                    coordinate = get_column_letter(c + 1) + str(r + 1)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                    merged = next(([a + 1, b, x + 1, y] for a, b, x, y in sheet.merged_cells
                                   if a <= r < b and x <= c < y), None)
                    text = coordinate + ": " + str(value)
                    cells.append(text)
                    blocks.append(block(text, "cell", sheet=sheet.name, row=r + 1,
                                        column=c + 1, cell=coordinate, merged_range=merged))
                if cells:
                    parts.append("\n".join(cells))
    finally:
        book.release_resources()
    return output(parts, blocks, findings=["LEGACY_XLS_VALUES_ONLY"],
                  limitations="旧版 XLS 保留单元格值与合并区域；公式表达式和内嵌图片需转换后复核。")


def parse_pptx(content):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    presentation = Presentation(io.BytesIO(content))
    parts, blocks, images, findings = [], [], {}, []
    if len(presentation.slides) > 1000:
        raise ValueError("SLIDE_LIMIT")

    def shapes(items, slide_number):
        for index, shape in enumerate(items):
            loc = {"slide": slide_number, "shape_id": shape.shape_id, "shape_index": index}
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                shapes(shape.shapes, slide_number)
            if shape.has_text_frame and shape.text.strip():
                parts.append(shape.text)
                blocks.append(block(shape.text, "slide_text", **loc))
            if shape.has_table:
                for r, row in enumerate(shape.table.rows):
                    values = [c.text for c in row.cells if not c.is_spanned]
                    text = json.dumps(values, ensure_ascii=False)
                    parts.append(text)
                    blocks.append(block(text, "table_row", **loc, row=r + 1))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    text = picture(shape.image.blob, images)
                    parts.append(text)
                    blocks.append(block(text, "image", **loc))
                except (ValueError, OSError):
                    findings.append("EMBEDDED_IMAGE_UNSUPPORTED")
            if shape.has_chart or shape.shape_type in {MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
                                                       MSO_SHAPE_TYPE.LINKED_OLE_OBJECT}:
                findings.append("EMBEDDED_OBJECT_REQUIRES_REVIEW")

    for number, slide in enumerate(presentation.slides, 1):
        header = "## 幻灯片 " + str(number)
        parts.append(header)
        blocks.append(block(header, "slide", slide=number))
        shapes(slide.shapes, number)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                text = "讲者备注：\n" + notes
                parts.append(text)
                blocks.append(block(text, "speaker_notes", slide=number))
    return output(parts, blocks, images=images, findings=sorted(set(findings)))


def parse_epub(content):
    from docreader.parser.epub_parser import EPUBParser
    parser = EPUBParser(file_name="source.epub")
    parts, blocks, images = [], [], {}
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if "META-INF/encryption.xml" in archive.namelist():
            raise ValueError("ENCRYPTED_EPUB_UNSUPPORTED")
        container = ET.fromstring(archive.read("META-INF/container.xml"))
        rootfile = container.find(".//{*}rootfile")
        if rootfile is None:
            raise ValueError("EPUB_PACKAGE_MISSING")
        package_path = rootfile.attrib["full-path"]
        package = ET.fromstring(archive.read(package_path))
        base = posixpath.dirname(package_path)
        manifest = {i.attrib["id"]: i.attrib for i in package.findall(".//{*}manifest/{*}item")}
        spine = package.findall(".//{*}spine/{*}itemref")
        if not spine or len(spine) > 1000:
            raise ValueError("EPUB_SPINE_INVALID")
        aliases = {}
        for item in manifest.values():
            if item.get("media-type", "").startswith("image/"):
                path = posixpath.normpath(posixpath.join(base, unquote(item["href"])))
                rendered = picture(archive.read(path), images)
                aliases[path] = rendered.split("](", 1)[1][:-1]
        for index, entry in enumerate(spine, 1):
            item = manifest[entry.attrib["idref"]]
            path = posixpath.normpath(posixpath.join(base, unquote(item["href"])))
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(archive.read(path), "lxml")
            for image in soup.find_all("img"):
                src = posixpath.normpath(posixpath.join(posixpath.dirname(path),
                                                      unquote(image.get("src", ""))))
                if src in aliases:
                    image["src"] = aliases[src]
            text = parser._html_to_markdown(str(soup))
            parts.append(text)
            blocks.append(block(text, "chapter", chapter=index, archive_path=path))
    return output(parts, blocks, images=images, reading_order="OPF spine")


def parse_xmind(content):
    from docreader.parser.xmind_parser import (
        XMindParser,
        _parse_json_sheets,
        _parse_xml_sheets,
        _read_content_entry,
    )
    document = XMindParser(file_name="source.xmind").parse(content)
    kind, payload = _read_content_entry(content)
    sheets = _parse_json_sheets(payload) if kind == "json" else _parse_xml_sheets(payload)
    blocks = []

    def walk(topic, sheet, path):
        if len(path) > 64 or len(blocks) > 10000:
            raise ValueError("XMIND_STRUCTURE_LIMIT")
        text = topic.title + ("\n" + topic.note if topic.note else "")
        blocks.append(block(text, "topic", sheet=sheet, topic_path=path))
        for index, child in enumerate(topic.children):
            walk(child, sheet, path + [index])

    for index, sheet in enumerate(sheets):
        if sheet.root_topic:
            walk(sheet.root_topic, index + 1, [0])
    findings = []
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if any(n.startswith(("attachments/", "resources/")) for n in archive.namelist()):
            findings.append("XMIND_ATTACHMENTS_REQUIRE_REVIEW")
    if b'"detached"' in payload or b'"relationships"' in payload:
        findings.append("XMIND_NON_TREE_CONTENT_REQUIRES_REVIEW")
    return output([document.content], blocks, findings=findings, **document.metadata)


def parse_web_archive(content, extension):
    from bs4 import BeautifulSoup
    from docreader.parser.html_parser import HTMLParser
    from docreader.parser.mhtml_parser import MHTMLParser
    parser = HTMLParser(file_name="source.html") if extension in {"html", "htm"} else MHTMLParser(file_name="source.mhtml")
    document = parser.parse(content)
    blocks = []
    if extension in {"html", "htm"}:
        soup = BeautifulSoup(content, "lxml")
        for unsafe in soup(["script", "style", "iframe", "object", "embed", "noscript"]):
            unsafe.decompose()
        for index, element in enumerate(soup.find_all(["h1", "h2", "h3", "p", "tr", "li"])):
            blocks.append(block(element.get_text(" ", strip=True), "html_element",
                                tag=element.name, element_index=index, element_id=element.get("id")))
    else:
        import email
        for index, part in enumerate(email.message_from_bytes(content).walk()):
            if part.get_content_type() == "text/html":
                soup = BeautifulSoup(part.get_payload(decode=True) or b"", "lxml")
                for unsafe in soup(["script", "style", "iframe", "object", "embed", "noscript"]):
                    unsafe.decompose()
                blocks.append(block(soup.get_text(" ", strip=True), "mime_html", mime_part=index,
                                    content_location=str(part.get("Content-Location", ""))))
    return output([document.content], blocks, images=document.images, **document.metadata)


PARSERS = {"json": parse_json, "xlsx": parse_xlsx, "xls": parse_xls,
           "pptx": parse_pptx, "epub": parse_epub, "xmind": parse_xmind}


def parse_format(content, extension):
    if extension in {"html", "htm", "mhtml", "mht"}:
        return parse_web_archive(content, extension)
    return PARSERS[extension](content)
