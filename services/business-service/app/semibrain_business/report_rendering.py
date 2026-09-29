"""Deterministic Markdown -> DOCX/PDF, with no network or model rewrite."""

import base64
import io
from uuid import uuid4

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt
from PIL import Image

from semibrain_business.sandbox import provider

VERSION = "report-markdown-v3"


def inline(paragraph, children, images, width=5.7):
    bold = italic = False
    link = None
    for token in children or []:
        if token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type == "link_open":
            link = token.attrGet("href")
        elif token.type == "link_close":
            if link and link.startswith(("https://", "http://")):
                paragraph.add_run(" (" + link + ")")
            link = None
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif token.type == "image":
            picture = images.get(token.attrGet("src"))
            if picture:
                with Image.open(io.BytesIO(picture)) as decoded:
                    fitted_width = min(width, 6.5 * decoded.width / decoded.height)
                paragraph.add_run().add_picture(io.BytesIO(picture), width=Inches(fitted_width))
                if token.content:
                    paragraph.add_run("\n" + token.content)
            else:
                paragraph.add_run("[图片：" + (token.content or "未登记图片") + "]")
        elif token.type in {"text", "code_inline", "html_inline"}:
            run = paragraph.add_run(token.content)
            run.bold, run.italic = bold, italic
            if token.type == "code_inline":
                run.font.name = "DejaVu Sans Mono"


def word(body, citations, images):
    if len(body) > 64000:
        raise ValueError("EXPORT_BODY_TOO_LARGE")
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    tokens = MarkdownIt("commonmark", {"html": False}).enable("table").parse(body)
    paragraph = None
    lists, quote_depth, table, cells = [], 0, None, []
    for token in tokens:
        kind = token.type
        if kind == "heading_open":
            paragraph = doc.add_heading(level=min(6, int(token.tag[1:])))
        elif kind == "bullet_list_open":
            lists.append({"ordered": False, "next": 1})
        elif kind == "ordered_list_open":
            lists.append({"ordered": True, "next": int(token.attrGet("start") or 1)})
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif kind == "list_item_open":
            parent = lists[-1]
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2 * len(lists))
            paragraph.add_run((str(parent["next"]) + ". ") if parent["ordered"] else "• ")
            parent["next"] += 1
        elif kind == "blockquote_open":
            quote_depth += 1
        elif kind == "blockquote_close":
            quote_depth -= 1
        elif kind == "paragraph_open":
            if not lists or paragraph is None:
                paragraph = doc.add_paragraph()
            if quote_depth:
                paragraph.paragraph_format.left_indent = Inches(0.25)
        elif kind in {"paragraph_close", "heading_close", "list_item_close"}:
            paragraph = None
        elif kind == "table_open":
            table = doc.add_table(rows=0, cols=0)
            table.style = "Table Grid"
        elif kind == "tr_open":
            cells = []
        elif kind in {"th_open", "td_open"}:
            cells.append([])
        elif kind == "inline":
            if table is not None:
                if cells:
                    cells[-1] = token.children or []
            else:
                if paragraph is None:
                    paragraph = doc.add_paragraph()
                inline(paragraph, token.children, images)
        elif kind == "tr_close":
            while len(table.columns) < len(cells):
                table.add_column(Inches(6.8 / max(1, len(cells))))
            row = table.add_row()
            for index, content in enumerate(cells):
                inline(
                    row.cells[index].paragraphs[0], content, images, width=6.3 / max(1, len(cells))
                )
        elif kind == "table_close":
            table = None
        elif kind in {"fence", "code_block"}:
            code = doc.add_paragraph()
            run = code.add_run(token.content.rstrip("\n"))
            run.font.name, run.font.size = "DejaVu Sans Mono", Pt(9)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "F3F5F4")
            code._p.get_or_add_pPr().append(shade)
        elif kind == "hr":
            doc.add_paragraph("—" * 30)
    if citations:
        doc.add_heading("来源", level=2)
        for source in citations:
            doc.add_paragraph(
                "[" + str(source.get("marker", "")) + "] " + str(source.get("title") or "来源")
            )
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def pdf(docx):
    identity = uuid4().hex + uuid4().hex
    sandbox = provider()
    code = """from pathlib import Path
import subprocess
subprocess.run(['/usr/bin/soffice','-env:UserInstallation=file:///tmp/export-profile','--headless','--nologo','--nodefault','--norestore','--convert-to','pdf','--outdir','/workspace','/workspace/report.docx'],check=True,timeout=24,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
assert Path('/workspace/report.pdf').is_file()
"""
    try:
        result = sandbox.execute(
            {
                "identity": identity,
                "code": code,
                "seconds": 28,
                "files": {"report.docx": base64.b64encode(docx).decode()},
                "exports": ["report.pdf"],
            }
        )
        if result.get("exit_code") != 0 or len(result.get("files", [])) != 1:
            raise ValueError("EXPORT_PDF_CONVERSION_FAILED")
        content = base64.b64decode(result["files"][0]["base64"], validate=True)
        if not content.startswith(b"%PDF-"):
            raise ValueError("EXPORT_PDF_INVALID")
        return content
    finally:
        sandbox.destroy(identity)
