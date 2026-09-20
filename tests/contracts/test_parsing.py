import base64
import hashlib
import io
import threading
import zipfile

import pytest
from docx import Document
from PIL import Image
from reportlab.pdfgen import canvas
from semibrain_business.parsing import ParseProfile, parse_asset, validate_archive


def test_markdown_table_and_original_locations(tmp_path):
    path = tmp_path / "sample.md"
    path.write_text(
        "# Process guide\n\n| Metric | Value |\n|---|---|\n| Yield | 91.5% |\n\nUnrelated text.\n",
        encoding="utf-8",
    )
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    result = parse_asset(path, allowed_root=tmp_path)
    assert result.status == "staged"
    assert "91.5%" in result.markdown
    assert any(b.location.get("line_start") == 1 for b in result.blocks)
    assert result.source_hash == before == hashlib.sha256(path.read_bytes()).hexdigest()


def test_docx_merge_values_and_native_fallback(tmp_path):
    document = Document()
    document.add_heading("Yield investigation", 1)
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Lot"
    table.cell(0, 1).text = "Yield"
    table.cell(1, 0).text = "Alpha"
    table.cell(2, 0).text = ""
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 1).text = "91.5%"
    table.cell(2, 1).text = "88.2%"
    path = tmp_path / "merge.docx"
    document.save(path)
    result = parse_asset(path, allowed_root=tmp_path)
    assert result.status == "staged"
    assert all(v in result.markdown for v in ["Yield investigation", "Alpha", "91.5%", "88.2%"])
    assert any(b.kind == "table_row" and b.location.get("row") == 2 for b in result.blocks)
    fallback = parse_asset(
        path,
        allowed_root=tmp_path,
        profile=ParseProfile(engines=("unavailable_engine", "weknora_docx")),
    )
    assert fallback.status == "staged"
    assert fallback.attempts[0].code == "ENGINE_UNAVAILABLE"
    assert fallback.attempts[1].status == "succeeded"


def test_pdf_actual_page_numbers_and_scanned_page(tmp_path):
    path = tmp_path / "digital.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 700, "Yield report 91.5 percent")
    pdf.showPage()
    pdf.drawString(50, 700, "Second page 88.2 percent")
    pdf.save()
    result = parse_asset(path, allowed_root=tmp_path)
    assert result.status == "staged"
    assert "91.5" in result.markdown and "88.2" in result.markdown
    assert {b.location.get("page_index") for b in result.blocks} == {0, 1}
    assert result.attempts[0].status == "skipped"
    image_path = tmp_path / "scan.png"
    Image.new("RGB", (700, 900), "grey").save(image_path)
    scan = tmp_path / "scan.pdf"
    pdf = canvas.Canvas(str(scan), pagesize=(700, 900))
    pdf.drawImage(str(image_path), 0, 0, 700, 900)
    pdf.save()
    result = parse_asset(scan, allowed_root=tmp_path)
    assert result.status == "needs_attention"
    assert result.images and "OCR_REQUIRED" in result.quality_findings


def test_csv_encoding_quoted_newline_and_formula_are_data(tmp_path):
    path = tmp_path / "data.csv"
    path.write_bytes(
        '批次;备注;数值\nLot-a;"first\nsecond";=1+1\nLot-A;普通;2.5\n'.encode("gb18030")
    )
    result = parse_asset(path, allowed_root=tmp_path)
    assert result.status == "staged"
    assert "Lot-a" in result.markdown and "Lot-A" in result.markdown
    assert "=1+1" in result.markdown and "2.5" in result.markdown
    assert result.blocks[1].location["line_start"] == 2
    assert result.blocks[1].location["line_end"] == 3


def test_all_failed_and_image_only_are_not_publishable(tmp_path):
    empty = tmp_path / "empty.md"
    empty.write_text("")
    result = parse_asset(
        empty,
        allowed_root=tmp_path,
        profile=ParseProfile(engines=("unavailable_engine", "weknora_markdown")),
    )
    assert result.status == "failed" and len(result.attempts) == 2
    image = io.BytesIO()
    Image.new("RGB", (80, 80), "red").save(image, format="PNG")
    path = tmp_path / "only-image.md"
    path.write_text(
        "![](data:image/png;base64," + base64.b64encode(image.getvalue()).decode() + ")"
    )
    result = parse_asset(path, allowed_root=tmp_path)
    assert result.status == "needs_attention" and result.images


def test_worker_cancel_timeout_and_path_boundary(tmp_path):
    path = tmp_path / "normal.md"
    path.write_text("# Content")
    cancelled = threading.Event()
    cancelled.set()
    assert parse_asset(path, allowed_root=tmp_path, cancel=cancelled).status == "cancelled"
    result = parse_asset(path, allowed_root=tmp_path, profile=ParseProfile(timeout_seconds=0.001))
    assert result.status == "failed" and "DEADLINE_EXCEEDED" in result.quality_findings
    during = threading.Event()
    timer = threading.Timer(0.01, during.set)
    timer.start()
    try:
        assert parse_asset(path, allowed_root=tmp_path, cancel=during).status == "cancelled"
    finally:
        timer.join()
    other = tmp_path / "other"
    other.mkdir()
    with pytest.raises(PermissionError):
        parse_asset(path, allowed_root=other)


def test_archive_traversal_and_corrupt_document(tmp_path):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("../outside", "payload")
    with pytest.raises(ValueError, match="UNSAFE_ARCHIVE_PATH"):
        validate_archive(buffer.getvalue())
    path = tmp_path / "corrupt.docx"
    path.write_bytes(b"not a document")
    assert parse_asset(path, allowed_root=tmp_path).status == "failed"
