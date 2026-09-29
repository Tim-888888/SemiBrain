import base64
import io
from datetime import timedelta
from types import SimpleNamespace
from unittest.mock import Mock
from uuid import uuid4

import pytest
from docx import Document
from fastapi import HTTPException
from PIL import Image
from pydantic import ValidationError
from semibrain_business import knowledge_api, report_exports, report_rendering
from semibrain_common.runtime import digest, now
from semibrain_contracts.exports import ExportRequest, ExportSource
from semibrain_conversation import exports as gateway


def picture():
    out = io.BytesIO()
    Image.new("RGB", (120, 60), "green").save(out, format="PNG")
    return out.getvalue()


@pytest.mark.parametrize(
    "data", [{"body_markdown": "forged"}, {"owner_id": str(uuid4())}, {"format": "html"}]
)
def test_exports_never_accept_client_body_owner_or_executable_format(data):
    with pytest.raises(ValidationError):
        ExportRequest(**{"request_id": uuid4(), "format": "docx", **data})


def test_word_keeps_chinese_tables_code_citations_and_registered_picture():
    body = "# 量测结果\n\n正文**粗体**和`CD`。\n\n|阶段|数值|\n|---|---|\n|CP|3.5|\n\n1. 首测\n2. 复测\n\n> 未证明因果\n\n```python\nprint('原样代码')\n```\n\n![量测示意](/registered)\n\n![未知图](https://example.org/no-fetch.png)\n\n尾句保持。"
    raw = report_rendering.word(
        body, [{"marker": "1", "title": "公开量测说明"}], {"/registered": picture()}
    )
    doc = Document(io.BytesIO(raw))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert all(
        part in text
        for part in [
            "量测结果",
            "正文粗体和CD",
            "1. 首测",
            "2. 复测",
            "print('原样代码')",
            "尾句保持。",
            "[1] 公开量测说明",
            "未知图",
        ]
    )
    assert doc.tables[0].cell(1, 0).text == "CP" and doc.tables[0].cell(1, 1).text == "3.5"
    assert len(doc.inline_shapes) == 1
    assert any(run.bold for p in doc.paragraphs for run in p.runs if run.text == "粗体")


def test_only_used_registered_images_are_reauthorized(monkeypatch):
    read = Mock(return_value=SimpleNamespace(body=picture()))
    monkeypatch.setattr(knowledge_api, "asset_response", read)
    result = report_exports.report_images(
        {
            "body_markdown": "![图](/one)\n![外部](https://unknown/img)",
            "claim": {"subject_id": "u"},
            "image_refs": [
                {"asset_id": "first", "url": "/one"},
                {"asset_id": "unused", "url": "/two"},
            ],
        }
    )
    assert list(result) == ["/one"]
    read.assert_called_once_with("first", {"subject_id": "u"})


@pytest.mark.parametrize(
    "result",
    [
        {"exit_code": 1, "files": []},
        {"exit_code": 0, "files": [{"base64": base64.b64encode(b"not pdf").decode()}]},
    ],
)
def test_pdf_failure_still_destroys_isolated_sandbox(monkeypatch, result):
    sandbox = Mock()
    sandbox.execute.return_value = result
    monkeypatch.setattr(report_rendering, "provider", lambda: sandbox)
    with pytest.raises(ValueError):
        report_rendering.pdf(b"docx")
    sandbox.destroy.assert_called_once_with(sandbox.execute.call_args.args[0]["identity"])


@pytest.mark.parametrize(
    "claim,expiry,code",
    [
        ({"subject_id": "other", "auth_version": 1}, 1, "EXPORT_UNAVAILABLE"),
        ({"subject_id": "owner", "auth_version": 2}, 1, "EXPORT_UNAVAILABLE"),
        ({"subject_id": "owner", "auth_version": 1}, -1, "EXPORT_EXPIRED"),
    ],
)
def test_export_download_checks_owner_epoch_and_expiry(monkeypatch, claim, expiry, code):
    call = Mock()
    monkeypatch.setattr(report_exports, "call", call)
    with pytest.raises(HTTPException) as failure:
        report_exports.source(
            {"owner_id": "owner", "auth_version": 1, "expires_at": now() + timedelta(days=expiry)},
            claim,
        )
    assert failure.value.detail["code"] == code
    call.assert_not_called()


@pytest.mark.parametrize(
    "report",
    [{"status": "running"}, {"status": "succeeded", "report_id": "r", "content_hash": "0" * 64}],
)
def test_gateway_rejects_nonfinal_or_changed_report(monkeypatch, report):
    monkeypatch.setattr(gateway, "internal_identity", lambda *a: None)
    user = {"_id": str(uuid4()), "auth_version": 1, "role": "user"}
    store = SimpleNamespace(users=Mock())
    store.users.find_one.return_value = user
    monkeypatch.setattr(gateway, "db", lambda: store)
    monkeypatch.setattr(gateway, "run_snapshot", lambda *a: report)
    with pytest.raises(HTTPException):
        gateway.source(
            ExportSource(
                run_id=uuid4(), owner_id=user["_id"], auth_version=1, content_hash="f" * 64
            ),
            None,
        )


def job_environment(monkeypatch, *, revoked=False, winning=True):
    body = "# 原始回答\n\n不改写的中文、**重点**与 [1]。\n"
    row = {
        "_id": "export",
        "run_id": "run",
        "owner_id": "owner",
        "auth_version": 1,
        "format": "md",
        "fence": 1,
        "content_hash": digest(body),
        "renderer_version": report_rendering.VERSION,
        "expires_at": now() + timedelta(days=1),
    }
    store = SimpleNamespace(report_exports=Mock(), assets=Mock())
    store.report_exports.update_one.return_value = SimpleNamespace(
        modified_count=1 if winning else 0
    )
    monkeypatch.setattr(report_exports, "db", lambda: store)
    monkeypatch.setattr(report_exports, "admission", Mock(return_value=row))
    source = Mock(return_value={"body_markdown": body})
    if revoked:
        source.side_effect = [{"body_markdown": body}, HTTPException(403)]
    monkeypatch.setattr(report_exports, "source", source)
    upload = Mock(return_value={"_id": "asset", "object_key": "owner/key"})
    monkeypatch.setattr(report_exports, "store_asset", upload)
    objects = Mock()
    monkeypatch.setattr(report_exports, "objects", lambda: objects)
    return body, store, upload, objects


def test_export_preserves_exact_markdown_without_llm(monkeypatch):
    body, store, upload, _ = job_environment(monkeypatch)
    assert report_exports.process_one()
    assert upload.call_args.args[0] == body.encode("utf-8")
    assert upload.call_args.kwargs["report_export_id"] == "export"
    update = store.report_exports.update_one.call_args
    assert update.args[0]["fence"] == 1 and update.args[1]["$set"]["status"] == "succeeded"


def test_revocation_during_conversion_prevents_upload(monkeypatch):
    _, store, upload, _ = job_environment(monkeypatch, revoked=True)
    report_exports.process_one()
    upload.assert_not_called()
    assert store.report_exports.update_one.call_args.args[1]["$set"]["status"] == "failed"


def test_losing_worker_cannot_publish_or_leave_an_asset(monkeypatch):
    _, store, _, objects = job_environment(monkeypatch, winning=False)
    report_exports.process_one()
    objects.remove_object.assert_called_once()
    store.assets.delete_one.assert_called_once_with({"_id": "asset", "owner_id": "owner"})


def test_expired_exports_remove_objects_before_metadata(monkeypatch):
    store = SimpleNamespace(report_exports=Mock(), assets=Mock())
    store.assets.find.side_effect = [
        Mock(limit=lambda n: []),
        [{"_id": "asset", "object_key": "key"}],
    ]
    store.report_exports.find.return_value.limit.return_value = [{"_id": "export"}]
    monkeypatch.setattr(report_exports, "db", lambda: store)
    objects = Mock()
    monkeypatch.setattr(report_exports, "objects", lambda: objects)
    report_exports.sweep()
    objects.remove_object.assert_called_once()
    store.assets.delete_one.assert_called_once()
    assert "purged_at" in store.report_exports.update_one.call_args.args[1]["$set"]
