import io
import json
import zipfile

import pytest
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from semibrain_business.format_parsers import parse_format
from semibrain_business.parsing import parse_asset, validate_archive


def archive(entries):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for path, value in entries.items():
            z.writestr(path, value)
    return out.getvalue()


def test_spreadsheet_retains_formula_cache_status_merges_and_original_coordinates(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Process statistics"
    sheet["A1"] = "Batch"
    sheet["B2"] = 47.25
    sheet["A2"] = "Synthetic-B"
    sheet.merge_cells("A2:A3")
    sheet["C2"] = "=B2*2"
    path = tmp_path / "coordinates.xlsx"
    book.save(path)
    result = parse_asset(path, allowed_root=tmp_path)
    assert result.status == "needs_attention"
    assert "47.25" in result.markdown and "=B2*2" in result.markdown
    assert "FORMULA_CACHE_MISSING" in result.quality_findings
    assert any(b.location.get("merged_range") == "A2:A3" for b in result.blocks)
    assert any(b.location.get("cell") == "C2" and b.location["cached_value"] is None
               for b in result.blocks)


def test_presentation_preserves_notes_image_and_slide_number(tmp_path):
    image = io.BytesIO()
    Image.new("RGB", (80, 80), "blue").save(image, format="PNG")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Probe contact"
    slide.notes_slide.notes_text_frame.text = "Review contact resistance before retesting."
    slide.shapes.add_picture(io.BytesIO(image.getvalue()), Inches(1), Inches(2))
    path = tmp_path / "review.pptx"
    deck.save(path)
    result = parse_asset(path, allowed_root=tmp_path)
    assert result.status == "staged" and result.images
    assert "Probe contact" in result.markdown and "resistance" in result.markdown
    assert all(b.location["slide"] == 1 for b in result.blocks)
    assert any(b.kind == "speaker_notes" for b in result.blocks)


def test_json_pointer_duplicate_keys_and_depth():
    result = parse_format('{"metrics/a":[{"~lot":"D2","yield":0.832}]}'.encode(), "json")
    assert {b["location"]["json_pointer"] for b in result["blocks"]} == {
        "/metrics~1a/0/~0lot", "/metrics~1a/0/yield"}
    assert "0.832" in result["content"]
    with pytest.raises(ValueError, match="JSON_DUPLICATE_KEY"):
        parse_format(b'{"n":1,"n":2}', "json")
    with pytest.raises(ValueError, match="JSON_NONFINITE"):
        parse_format(b'{"n":NaN}', "json")
    with pytest.raises(ValueError, match="JSON_STRUCTURE_LIMIT"):
        parse_format(b"[" * 70 + b"0" + b"]" * 70, "json")


def test_xmind_json_xml_notes_and_attachment_degradation():
    obj = [{"title": "Investigation", "rootTopic": {"title": "Defect", "children": {
        "attached": [{"title": "Contact", "notes": {"plain": {"content": "Check fixture"}}}]}}}]
    data = archive({"content.json": json.dumps(obj), "attachments/chart.png": b"uninterpreted"})
    result = parse_format(data, "xmind")
    assert "Check fixture" in result["content"]
    assert result["blocks"][1]["location"]["topic_path"] == [0, 0]
    assert result["metadata"]["quality_findings"] == ["XMIND_ATTACHMENTS_REQUIRE_REVIEW"]
    xml = '<xmap-content><sheet><title>Flow</title><topic><title>Alpha</title><notes><plain>Beta</plain></notes></topic></sheet></xmap-content>'
    result = parse_format(archive({"content.xml": xml}), "xmind")
    assert "Alpha" in result["content"] and "Beta" in result["content"]


def test_epub_spine_order_overrides_zip_order_and_same_basename_images_stay_distinct():
    images = []
    for color in ["blue", "red"]:
        buf = io.BytesIO()
        Image.new("RGB", (64, 64), color).save(buf, format="PNG")
        images.append(buf.getvalue())
    entries = {
        "META-INF/container.xml": '<container><rootfiles><rootfile full-path="book/content.opf"/></rootfiles></container>',
        "book/content.opf": '<package><manifest><item id="b" href="b/page.xhtml" media-type="application/xhtml+xml"/><item id="a" href="a/page.xhtml" media-type="application/xhtml+xml"/><item id="ia" href="a/chart.png" media-type="image/png"/><item id="ib" href="b/chart.png" media-type="image/png"/></manifest><spine><itemref idref="a"/><itemref idref="b"/></spine></package>',
        "book/b/page.xhtml": '<html><body><h1>Second process</h1><img src="chart.png"/></body></html>',
        "book/a/page.xhtml": '<html><body><h1>First process</h1><img src="chart.png"/></body></html>',
        "book/a/chart.png": images[0], "book/b/chart.png": images[1],
    }
    result = parse_format(archive(entries), "epub")
    assert result["content"].index("First process") < result["content"].index("Second process")
    assert len(result["images"]) == 2
    keys = list(result["images"])
    assert keys[0] in result["blocks"][0]["text"] and keys[1] in result["blocks"][1]["text"]
    with pytest.raises(ValueError, match="ENCRYPTED_EPUB"):
        parse_format(archive({**entries, "META-INF/encryption.xml": "<encryption/>"}), "epub")


def test_html_encoding_no_script_and_dom_locator(tmp_path):
    path = tmp_path / "sample.html"
    path.write_bytes('<html><head><meta charset="gb18030"></head><body><h1>晶圆检测</h1><p id="result">检测值：31.7</p><script>alert("ignored")</script></body></html>'.encode("gb18030"))
    result = parse_asset(path, allowed_root=tmp_path)
    assert result.status == "staged"
    assert "晶圆检测" in result.markdown and "31.7" in result.markdown
    assert "alert" not in result.markdown
    assert any(b.location.get("element_id") == "result" for b in result.blocks)


def test_macro_container_not_executed_or_published():
    with pytest.raises(ValueError, match="OFFICE_MACROS_UNSUPPORTED"):
        validate_archive(archive({"xl/vbaProject.bin": b"not-executable"}))


def test_image_without_cloud_consent_keeps_original_but_requires_review(tmp_path, monkeypatch):
    monkeypatch.setenv("SEMIBRAIN_VISION_API_KEY", "not-used")
    monkeypatch.setenv("SEMIBRAIN_VISION_BASE_URL", "https://must-not-call.invalid")
    path = tmp_path / "drawing.png"
    Image.new("RGB", (96, 96), "gray").save(path)
    result = parse_asset(path, allowed_root=tmp_path)
    assert result.status == "needs_attention" and result.images
    assert result.attempts[0].status == "skipped"
    assert result.attempts[0].code == "EXTERNAL_DATA_DENIED"


def test_human_review_retains_original_snapshot_and_rejects_new_assets():
    from uuid import uuid4

    from fastapi import HTTPException
    from semibrain_business.knowledge_revisions import reviewed_snapshot
    from semibrain_business.parsing import ParseResult
    from semibrain_contracts.knowledge_review import ReviewParsed

    original = ParseResult(status="needs_attention", source_hash="a" * 64, markdown="unclear",
                           quality_findings=["FORMULA_CACHE_MISSING"])
    form = ReviewParsed(request_id=uuid4(), expected_revision=1, source_version=uuid4(),
                        text="Checked value: 27.3; no computed formula result is claimed.",
                        reason="Manually compared against the source workbook.")
    result = reviewed_snapshot(original, form, str(uuid4()), "reviewer")
    assert result.status == "staged" and not result.quality_findings
    assert result.parser_manifest["human_review"]["findings"] == ["FORMULA_CACHE_MISSING"]
    assert original.markdown == "unclear" and original.status == "needs_attention"
    with pytest.raises(HTTPException):
        reviewed_snapshot(original, form.model_copy(update={"text": "![fake](/v1/assets/other/content)"}), str(uuid4()), "reviewer")


def test_ingestion_cancel_and_revocation_stop_parser(monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import Mock

    from semibrain_business import knowledge

    store = SimpleNamespace(ingestion_jobs=Mock(), documents=Mock())
    store.ingestion_jobs.find_one.return_value = {"cancel_requested": False}
    store.documents.find_one.return_value = {"revision": 4, "revoked": False}
    monkeypatch.setattr(knowledge, "db", lambda: store)
    job = {"_id": "job", "document_id": "doc", "generation": 4}
    assert not knowledge.IngestionCancellation(job, "fence").is_set()
    store.ingestion_jobs.find_one.return_value["cancel_requested"] = True
    assert knowledge.IngestionCancellation(job, "fence").is_set()
    store.ingestion_jobs.find_one.return_value["cancel_requested"] = False
    store.documents.find_one.return_value["revoked"] = True
    assert knowledge.IngestionCancellation(job, "fence").is_set()
